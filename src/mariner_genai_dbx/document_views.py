"""Read-only HTML views for documents the browser cannot display itself.

Spreadsheets are shown as one table per visible sheet, anchored ``#sheet-N``
(1-based, visible sheets only) so citations can open the cited sheet. Values
are the ones Excel last saved; formulas are never evaluated.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from html import escape
from io import BytesIO


MAX_ROWS = 5000
MAX_COLUMNS = 100

# The pages contain no scripts; styles are inline and every value is escaped.
CONTENT_SECURITY_POLICY = (
    "default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; "
    "form-action 'none'; frame-ancestors 'self'"
)


class SpreadsheetError(Exception):
    """The workbook could not be read."""


@dataclass(frozen=True)
class Sheet:
    name: str
    rows: tuple[tuple[int, tuple[str, ...]], ...]  # (Excel row number, cell texts)
    truncated: bool


def read_workbook(data: bytes, suffix: str) -> list[Sheet]:
    try:
        if suffix.lower() == ".xls":
            return _read_xls(data)
        return _read_xlsx(data)
    except SpreadsheetError:
        raise
    except Exception as error:
        raise SpreadsheetError("The workbook could not be read") from error


def _read_xlsx(data: bytes) -> list[Sheet]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return [
            _sheet(worksheet.title, worksheet.iter_rows(values_only=True))
            for worksheet in workbook.worksheets
            if worksheet.sheet_state == "visible"
        ]
    finally:
        workbook.close()


def _read_xls(data: bytes) -> list[Sheet]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)

    def value(cell):
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, workbook.datemode)
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return xlrd.error_text_from_code.get(cell.value, "#ERROR")
        return cell.value

    sheets = []
    for index in range(workbook.nsheets):
        worksheet = workbook.sheet_by_index(index)
        if worksheet.visibility == 0:
            rows = (
                [value(cell) for cell in worksheet.row(row)]
                for row in range(worksheet.nrows)
            )
            sheets.append(_sheet(worksheet.name, rows))
    workbook.release_resources()
    return sheets


def _sheet(name, raw_rows) -> Sheet:
    rows = []
    truncated = False
    for number, raw in enumerate(raw_rows, start=1):
        cells = [_text(value) for value in list(raw)[:MAX_COLUMNS]]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        if len(rows) == MAX_ROWS:
            truncated = True
            break
        rows.append((number, tuple(cells)))
    return Sheet(name=str(name), rows=tuple(rows), truncated=truncated)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def render_workbook(title: str, sheets: list[Sheet], download_url: str) -> str:
    if not any(sheet.rows for sheet in sheets):
        return render_message(
            title, "This workbook has no visible data to preview.", download_url
        )
    tabs = []
    sections = []
    active_rules = []
    for number, sheet in enumerate(sheets, start=1):
        anchor = f"sheet-{number}"
        tabs.append(f'<a href="#{anchor}" class="tab tab-{number}">{escape(sheet.name)}</a>')
        active_rules.append(
            f"body:has(#{anchor}:target) .tab-{number}"
            + (f", body:not(:has(.sheet:target)) .tab-{number}" if number == 1 else "")
            + " { color: #ffffff; background: #013474; }"
        )
        sections.append(_render_sheet(anchor, sheet, download_url))
    return _page(
        title,
        extra_style="\n".join(active_rules),
        body=f'<nav aria-label="Sheets">{"".join(tabs)}</nav>' + "".join(sections),
    )


def _render_sheet(anchor: str, sheet: Sheet, download_url: str) -> str:
    if not sheet.rows:
        return f'<section id="{anchor}" class="sheet"><p class="note">This sheet is empty.</p></section>'
    width = max(len(cells) for _number, cells in sheet.rows)
    header_number, header_cells = sheet.rows[0]
    head = "".join(
        f"<th>{escape(cell)}</th>" for cell in header_cells + ("",) * (width - len(header_cells))
    )
    body = "".join(
        f'<tr><th class="row-number">{number}</th>'
        + "".join(f"<td>{escape(cell)}</td>" for cell in cells)
        + "<td></td>" * (width - len(cells))
        + "</tr>"
        for number, cells in sheet.rows[1:]
    )
    note = (
        f'<p class="note">Showing the first {MAX_ROWS:,} rows. '
        f'<a href="{escape(download_url)}" download>Download</a> the file to see all of it.</p>'
        if sheet.truncated
        else ""
    )
    return (
        f'<section id="{anchor}" class="sheet">{note}<div class="table-wrap"><table>'
        f'<thead><tr><th class="row-number">{header_number}</th>{head}</tr></thead>'
        f"<tbody>{body}</tbody></table></div></section>"
    )


def render_message(title: str, message: str, download_url: str) -> str:
    return _page(
        title,
        body=(
            f'<p class="note">{escape(message)} '
            f'<a href="{escape(download_url)}" download>Download the file</a></p>'
        ),
    )


def _page(title: str, body: str, extra_style: str = "") -> str:
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title>
<style>
  body {{ margin: 0; color: #171d43; background: #ffffff;
    font: 14px/1.4 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; }}
  nav {{ display: flex; flex-wrap: wrap; gap: .35rem; padding: .6rem 1rem; border-bottom: 1px solid #dce2ea; }}
  .tab {{ padding: .3rem .65rem; color: #013474; background: #eef2f7; border-radius: 4px;
    font-weight: 600; text-decoration: none; }}
  .sheet {{ display: none; }}
  .sheet:target, body:not(:has(.sheet:target)) .sheet:first-of-type {{ display: block; }}
  .table-wrap {{ overflow: auto; max-height: calc(100vh - 3.5rem); }}
  table {{ border-collapse: collapse; font-size: 13px; }}
  th, td {{ padding: .3rem .55rem; border: 1px solid #dce2ea; text-align: left; vertical-align: top;
    white-space: pre-wrap; max-width: 28rem; }}
  thead th {{ position: sticky; top: 0; z-index: 1; background: #e8eef6; }}
  .row-number {{ color: #647184; background: #f6f8fb; font-weight: 400; text-align: right; }}
  .note {{ margin: .75rem 1rem; color: #647184; }}
{extra_style}
</style>
</head>
<body>{body}</body>
</html>"""
